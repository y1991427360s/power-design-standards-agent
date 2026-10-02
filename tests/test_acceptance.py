import io
import json
from docx import Document
from reportlab.pdfgen.canvas import Canvas
from parser.documents import parse
from agent.provider import LLMProvider

def upload(client, text, name='acceptance.txt', **meta):
    data={'standard_name':'自编功能验收资料（非工程标准）','standard_code':'TEST-001','version':'2026',
          'source_type':'STANDARD','standard_type':'其他','category':'电缆', **meta}
    raw=text.encode() if isinstance(text,str) else text
    return client.post('/api/documents',files={'file':(name,raw)},data={'metadata':json.dumps(data)})

def query(client, question):
    response=client.post('/api/chat',json={'question':question})
    assert response.status_code==200, response.text
    return response.json()

def test_clause_page_source_and_grounded_query(client):
    doc=upload(client,'1 总则\n1.0.1 本资料为自编功能验收资料。\f\n3 电缆\n3.2 控制电缆\n3.2.1 控制电缆屏蔽层应连接测试接地端子。').json()
    clauses=client.get('/api/documents/'+doc['document_id']+'/clauses').json()
    clause=next(r for r in clauses if r['clause']=='3.2.1')
    assert clause['page']==2
    assert clause['chapter']=='3 电缆'
    a=query(client,'控制电缆屏蔽层如何连接？')
    cited=next(r for r in a['citations'] if r['clause']=='3.2.1')
    assert cited['standard_code']=='TEST-001' and cited['version']=='2026' and cited['source_file']
    assert '控制电缆屏蔽层应连接测试接地端子。' in cited['text']
    assert client.get('/api/clauses/'+cited['id']).json()['context']
    assert client.get('/api/documents/'+doc['document_id']+'/file').status_code==200

def test_pdf_upload_physical_pages_and_cross_page_clause(client):
    file=io.BytesIO(); canvas=Canvas(file)
    canvas.drawString(50,750,'3.2.1 Cable shielding shall connect to earth terminal.')
    canvas.showPage(); canvas.drawString(50,750,'Continuation: verify continuity before energizing.')
    canvas.drawString(50,700,'3.2.2 Spare cores shall carry labels.')
    canvas.save()
    response=upload(client,file.getvalue(),'sample.pdf')
    assert response.status_code==200,response.text
    clauses=client.get('/api/documents/'+response.json()['document_id']+'/clauses').json()
    first=next(r for r in clauses if r['clause']=='3.2.1')
    assert first['page']==1 and first['page_end']==2 and first['page_kind']=='PDF物理页'
    assert 'Continuation' in first['text']
    assert next(r for r in clauses if r['clause']=='3.2.2')['page']==2

def test_docx_tables_and_no_invented_physical_page(client):
    doc=Document();doc.add_paragraph('3.2.1 控制电缆备用芯应编号。');table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='编号';table.cell(0,1).text='端子A';buf=io.BytesIO();doc.save(buf)
    response=upload(client,buf.getvalue(),'sample.docx')
    assert response.status_code==200,response.text
    row=client.get('/api/documents/'+response.json()['document_id']+'/clauses').json()[0]
    assert row['text'].count('3.2.1')==1 and '端子A' in row['text']
    assert '非排版页码' in row['page_kind']

def test_unknown_question_no_hallucinated_citation(client):
    upload(client,'3.2.1 控制电缆屏蔽层应连接测试接地端子。')
    a=query(client,'月球核聚变反应堆安全距离是多少？')
    assert a['conclusion']=='当前知识库未检索到明确规范依据。'
    assert a['citations']==[] and a['confidence']=='低'
    assert '能够直接支持该结论' in a['analysis']

def test_missing_requested_standard(client):
    upload(client,'3.2.1 控制电缆屏蔽层应连接测试接地端子。')
    a=query(client,'GB 999999 对控制电缆屏蔽层有什么要求？')
    assert not a['citations']

def test_multiple_standards_separate_sources(client):
    upload(client,'3.2.1 控制电缆屏蔽层应连接测试接地端子。',standard_code='TEST-A')
    upload(client,'3.2.1 控制电缆屏蔽层应记录连接位置。',standard_code='TEST-B')
    a=query(client,'控制电缆屏蔽层有什么要求？')
    assert {r['standard_code'] for r in a['citations']}=={'TEST-A','TEST-B'}
    assert '不自动判定' in a['analysis']

def test_experience_is_never_formal_evidence(client):
    upload(client,'1.0.1 紫色端子标记备用芯是某项目的历史经验。',source_type='PROJECT_EXPERIENCE')
    a=query(client,'紫色端子标记备用芯')
    assert a['citations']==[] and a['references'][0]['source_type']=='PROJECT_EXPERIENCE'
    assert a['conclusion']=='当前知识库未检索到明确规范依据。'

def test_demo_never_real_standard(client):
    upload(client,'DEMO / 非真实规范\n3.2.1 控制电缆屏蔽层应连接测试端子。')
    a=query(client,'控制电缆屏蔽层')
    assert a['citations'][0]['is_demo']==1 and a['confidence']=='低'
    assert a['conclusion'].startswith('当前知识库未检索到明确规范依据。')

def test_demo_metadata_and_heading_not_evidence(client):
    upload(client,'3 电缆\n3.2 控制电缆\n3.2.1 控制电缆屏蔽层应连接端子。',standard_code='DEMO-99')
    a=query(client,'控制电缆屏蔽层')
    assert all(r['is_demo'] and r['clause'] for r in a['citations'])

def test_reparse_version_and_delete(client):
    doc=upload(client,'3.2.1 控制电缆屏蔽层测试要求。').json();id=doc['document_id']
    assert client.post('/api/documents/'+id+'/reparse').json()['chunk_count']==1
    update={'standard_name':doc['standard_name'],'standard_code':'TEST-001','lifecycle':'废止'}
    assert client.patch('/api/documents/'+id,json=update).status_code==200
    assert query(client,'控制电缆屏蔽层')['citations']==[]
    assert client.delete('/api/documents/'+id).status_code==200
    assert client.get('/api/documents/'+id).status_code==404

def test_invalid_upload_and_scan_errors(client):
    assert upload(client,b'','empty.txt').status_code==400
    assert upload(client,'anything','sample.exe').status_code==400
    file=io.BytesIO();c=Canvas(file);c.showPage();c.save()
    assert upload(client,file.getvalue(),'scan.pdf').status_code==422
    assert client.get('/api/documents').json()[0]['status']=='解析失败'
    assert 'OCR' in client.get('/api/documents').json()[0]['warning']

def test_rules_and_history(client):
    r=client.post('/api/rules',json={'title':'测试判据','rule':'这是设计习惯','source':'自编'}).json()
    assert client.get('/api/rules').json()[0]['title']=='测试判据'
    assert client.put('/api/rules/'+r['id'],json={'title':'修订','rule':'新习惯'}).status_code==200
    a=query(client,'完全没有资料的问题')
    assert client.get('/api/conversations/'+a['conversation_id']).json()[0]['question']=='完全没有资料的问题'
    assert client.delete('/api/rules/'+r['id']).status_code==200

def test_external_disabled_and_cross_site_blocked(client, monkeypatch):
    monkeypatch.setenv('LLM_MODEL','configured-model')
    assert LLMProvider().select('test',[]) is None
    assert client.post('/api/chat',json={'question':'测试'},headers={'Origin':'https://evil.example'}).status_code==403
    assert client.get('/api/documents',headers={'Host':'evil.example'}).status_code==403
